import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_short_docx(output_path):
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    COLOR_PRIMARY = RGBColor(91, 33, 182)     # Deep Purple
    COLOR_DARK = RGBColor(30, 41, 59)         # Slate 800

    def style_run(run, font_name="TH Sarabun New", size_pt=16, bold=False, italic=False, color=COLOR_DARK):
        run.font.name = font_name
        run.font.size = Pt(size_pt)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
        rPr = run._r.get_or_add_rPr()
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" w:cs="{font_name}"/>')
        rPr.append(rFonts)

    def add_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        style_run(run, size_pt=22, bold=True, color=COLOR_PRIMARY)

    def add_subtitle(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(12)
        run = p.add_run(text)
        style_run(run, size_pt=15, italic=True, color=RGBColor(100, 116, 139))

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        style_run(run, size_pt=17, bold=True, color=COLOR_PRIMARY)

    def add_body(text, bold_prefix=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            style_run(rb, size_pt=15, bold=True, color=COLOR_DARK)
        run = p.add_run(text)
        style_run(run, size_pt=15, color=COLOR_DARK)

    def add_quote(text):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        cell = table.cell(0, 0)
        cell.width = Inches(6.5)
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F5F3FF"/>')
        tcPr.append(shd)
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="7C3AED"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
        tcPr.append(borders)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(f'"{text}"')
        style_run(run, size_pt=14, italic=True, color=COLOR_DARK)
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(0)
        p_sp.paragraph_format.space_after = Pt(3)

    add_title("บทพูดอธิบายโครงงานวิจัยให้อาจารย์เข้าใจง่าย (ฉบับย่อ)")
    add_subtitle("ระบบคลังรวมและเผยแพร่โครงงานวิจัย บัณฑิตศึกษาและโปรเจกต์จบ (มรภ.สุรินทร์)")

    add_h2("1. สรุปสั้นใน 30 วินาที (Elevator Pitch)")
    add_quote("อาจารย์ครับ โครงงานนี้คือ 'คลังงานวิจัยออนไลน์ของ มรภ.สุรินทร์' ที่มี AI ช่วยอำนวยความสะดวกครบวงจรครับ แทนที่นักศึกษาจะค้นหาได้แค่พิมพ์คำตรงเป๊ะ ระบบนี้ใช้ AI ช่วยแนะนำงานวิจัยที่เกี่ยวข้องกัน คอยตรวจว่าหัวข้อวิจัยที่กำลังจะทำซ้ำกับคนอื่นไหม และเปิดให้นักศึกษาแช็ตถาม-ตอบเนื้อหากับเล่ม PDF ได้โดยตรงเลยครับ")

    add_h2("2. คุยกับอาจารย์แบบ 3 ส่วนหลัก (ปัญหา -> ทางแก้ -> ประโยชน์)")
    add_body("นักศึกษาค้นหางานวิจัยเดิมยาก ต้องพิมพ์คำสะกดตรงเป๊ะถึงจะเจอ และบ่อยครั้งที่ทำหัวข้อวิจัยซ้ำกับรุ่นพี่ที่เคยทำไปแล้ว เพราะไม่มีระบบตรวจจับความซ้ำซ้อนส่วนกลางของมหาวิทยาลัยครับ", bold_prefix="• ปัญหาเดิม: ")
    add_body("เราจึงทำเว็บคลังรวมงานวิจัยทั้ง 7 คณะ โดยใส่ฟีเจอร์ AI เด่นๆ 4 อย่างเข้ามาช่วย:", bold_prefix="• สิ่งที่ระบบเราแก้ปัญหา: ")
    add_body("พิมพ์ค้นหาภาษาพูดธรรมดาได้ เช่น 'เกษตรกรรม' ระบบก็เข้าใจและดึงงาน 'ข้าวหอมมะลิ' มาให้", bold_prefix="   1. ค้นหาเชิงความหมาย (Semantic Search): ")
    add_body("พอนักศึกษาพิมพ์ชื่อหัวข้อ AI จะคำนวณทันทีว่าซ้ำกับงานเดิมกี่เปอร์เซ็นต์ พร้อมออกใบรับรองความเป็นงานวิจัยต้นฉบับได้", bold_prefix="   2. เรดาร์ตรวจจับความซ้ำซ้อน (Redundancy Radar): ")
    add_body("นักศึกษาเปิดเล่มวิจัยแล้วพิมพ์ถาม AI ได้เลย เช่น 'งานนี้ใช้วิธีวิจัยแบบไหน' ระบบจะดึงคำตอบพร้อมชี้เลขหน้าให้ทันที โดยไม่ต้องนั่งอ่านทั้งเล่ม", bold_prefix="   3. คุยกับไฟล์ PDF (AI Chat with PDF): ")
    add_body("แสดงแผนผังให้เห็นว่า คณะไหน อาจารย์ท่านใด เชี่ยวชาญงานวิจัยเรื่องอะไรบ้าง", bold_prefix="   4. แผนผังเครือข่ายความรู้ (Knowledge Graph): ")
    add_body("นักศึกษาทำวิจัยได้เร็วขึ้น ไม่ทำหัวข้อซ้ำซ้อน ส่วนอาจารย์ก็มีระบบตรวจตรวจแก้และไฮไลต์บน PDF ได้ในที่เดียว และที่สำคัญรันบนเครื่องเซิร์ฟเวอร์มหาวิทยาลัยได้เลย ไม่ต้องเสียค่าบริการ Cloud รายเดือนครับ", bold_prefix="• ประโยชน์ที่อาจารย์และมหาวิทยาลัยได้รับ: ")

    add_h2("3. จุดเด่นเวลาอาจารย์ถามเจาะลึกสั้นๆ")
    add_body("ระบบเราประมวลผลภาษาไทยด้วย PyThaiNLP และรันด้วยระบบของเครื่องเราเอง (On-Premise) ทำให้ไม่มีค่าใช้จ่าย API รายเดือน ข้อมูลงานวิจัยปลอดภัย และตอบสนองได้เร็วในเสี้ยววินาทีครับ", bold_prefix="• ถาม: 'ใช้ AI อะไร เปลืองเงินไหม?': ")
    add_body("ระบบนี้ไม่ได้ค้นหาแบบ LIKE ปกติ แต่คำนวณเวกเตอร์ความคล้ายคลึงของเนื้อหาผสมกับพฤติกรรมการเปิดอ่าน ซึ่งจากการทดสอบได้ค่าความแม่นยำ (F1-score) สูงถึง 0.924 ดีกว่าวิธีเดิม 29.7% ครับ", bold_prefix="• ถาม: 'ต่างจากเว็บค้นหาทั่วไปยังไง?': ")

    doc.save(output_path)
    print(f"Saved to {output_path}")

if __name__ == "__main__":
    create_short_docx("d:/SRRU-RESEARCH/docs/SHORT_EXPLANATION_FOR_ADVISOR.docx")
